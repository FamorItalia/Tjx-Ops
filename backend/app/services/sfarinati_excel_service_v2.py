from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
import re

from openpyxl import Workbook
from openpyxl.styles import Font

from app.schemas.packing_lists import PackingListRead


@dataclass
class SfarinatiExcelV2BuildResult:
    file_name: str
    file_path: str
    generated_at: datetime


class SfarinatiExcelServiceV2:
    def __init__(self, output_root: Path):
        self.output_root = output_root

    @staticmethod
    def _slug(value: str | None, fallback: str) -> str:
        if not value:
            return fallback
        return re.sub(r"[^A-Za-z0-9]+", "_", value.upper()).strip("_") or fallback

    def build_output_path(self, payload: PackingListRead) -> Path:
        out_dir = self.output_root / "sfarinati"
        out_dir.mkdir(parents=True, exist_ok=True)
        brand = self._slug(payload.brand, "BRAND")
        po = self._slug(payload.po_number, f"PO_{payload.id}")
        dc = self._slug(payload.dc_code, "GLOBAL")
        return out_dir / f"SFARINATI_{brand}_{po}_{dc}.xlsx"

    def build_excel(self, payload: PackingListRead) -> SfarinatiExcelV2BuildResult:
        output_path = self.build_output_path(payload)
        wb = Workbook()
        ws = wb.active
        ws.title = "SFARINATI"

        ws["A1"] = "OGGETTO: DICHIARAZIONE DI LIBERA ESPORTAZIONE"
        ws["A1"].font = Font(bold=True, size=13)
        ws["A2"] = "OGGETTO: SFARINATI E PASTE ALIMENTARI"
        ws["A2"].font = Font(bold=True, size=13)
        ws["A4"] = "Consapevoli di assumere ogni conseguente responsabilità, siamo a dichiararvi che le merci (NC.1902) esportate con"
        ws["A6"] = f"Numero fattura: {payload.invoice_number or '-'} - Data fattura: {payload.document_date or '-'}"
        ws["A6"].font = Font(bold=True)
        ws["A8"] = (
            "non rientrano nelle disposizioni previste dal Decreto del Ministero delle Politiche Agricole "
            "Alimentari e Forestali del 17 dicembre 2013, concernente l'obbligo di \"comunicazione\" prevista "
            "dall'art.12, comma 1, del Decreto del Presidente della Repubblica 9 febbraio 2001, n.187 "
            "(cod.09YY - prodotto non sottoposto a comunicazione)."
        )
        ws["A10"] = f"Livorno, {datetime.now().strftime('%d/%m/%Y')}"
        ws["A11"] = "Timbro e Firma"
        ws["A11"].font = Font(bold=True)
        ws.column_dimensions["A"].width = 140
        wb.save(output_path)
        return SfarinatiExcelV2BuildResult(
            file_name=output_path.name,
            file_path=str(output_path.resolve()),
            generated_at=datetime.fromtimestamp(output_path.stat().st_mtime),
        )
