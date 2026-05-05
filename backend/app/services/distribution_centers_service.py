from datetime import datetime, timezone
from io import BytesIO
import re

from openpyxl import Workbook, load_workbook
from sqlalchemy.orm import Session

from app.repositories.distribution_centers_repository import DistributionCentersRepository
from app.schemas.distribution_centers import DistributionCenterImportResponse


class DistributionCentersService:
    def __init__(self, db: Session):
        self.db = db
        self.repo = DistributionCentersRepository(db)

    @staticmethod
    def _normalize(value) -> str | None:
        if value is None:
            return None
        v = str(value).strip()
        return v if v else None

    @staticmethod
    def _normalize_brand(value: str | None) -> str | None:
        if value is None:
            return None
        raw = value.strip().upper()
        mapping = {
            "T.J.MAXX": "TJ MAXX",
            "TJ MAXX": "TJ MAXX",
            "MARSHALLS": "MARSHALLS",
            "HOME GOODS": "HOMEGOODS",
            "HOMEGOODS": "HOMEGOODS",
            "CAN MARSH": "CAN MARSH",
            "HOME SENSE": "HOME SENSE",
            "WINNERS": "WINNERS",
            "SIERRA": "SIERRA",
        }
        return mapping.get(raw, raw)

    @staticmethod
    def _extract_city_state_country(city_state: str | None, zip_row: str | None) -> tuple[str | None, str | None, str | None]:
        city = None
        state = None
        country = None
        if city_state:
            cleaned = city_state.strip()
            # typical "CITY, ST"
            m = re.match(r"^(.+?),\s*([A-Z]{2})$", cleaned)
            if m:
                city = m.group(1).strip()
                state = m.group(2).strip()
            else:
                # e.g. "L4V 1B8 MISSISSAUGA"
                m2 = re.match(r"^[A-Z0-9 ]+\s+([A-Z][A-Z ]+)$", cleaned)
                if m2:
                    city = m2.group(1).strip().title()
        if zip_row:
            z = zip_row.strip()
            if "CANADA" in z.upper():
                country = "CANADA"
            elif re.fullmatch(r"\d{5}", z):
                country = "USA"
        return city, state, country

    def import_distribution_centers_excel(self, file_name: str, content: bytes) -> DistributionCenterImportResponse:
        wb = load_workbook(BytesIO(content), data_only=True)
        ws = wb[wb.sheetnames[0]]

        rows = list(ws.iter_rows(values_only=True))
        if len(rows) < 6:
            raise ValueError("Formato Excel DC non valido: righe insufficienti.")

        brands_row = rows[0]
        po_prefix_row = rows[1]
        alt_dc_code_row = rows[2] if len(rows) > 2 else ()
        dc_code_row = rows[3]
        address_row = rows[4]
        city_state_row = rows[5]
        zip_row = rows[6] if len(rows) > 6 else ()
        general_division_row = rows[7] if len(rows) > 7 else ()
        general_addr1_row = rows[8] if len(rows) > 8 else ()
        general_addr2_row = rows[9] if len(rows) > 9 else ()
        general_addr3_row = rows[10] if len(rows) > 10 else ()
        dest_name_row = rows[11] if len(rows) > 11 else ()
        dest_dc_row = rows[12] if len(rows) > 12 else ()
        dest_addr1_row = rows[13] if len(rows) > 13 else ()
        dest_addr2_row = rows[14] if len(rows) > 14 else ()
        dest_addr3_row = rows[15] if len(rows) > 15 else ()

        inserted = 0
        updated = 0
        skipped = 0
        warnings: list[str] = []

        total_columns = 0
        max_cols = max(len(brands_row), len(po_prefix_row), len(dc_code_row), len(address_row), len(city_state_row))
        for col in range(1, max_cols):
            brand = self._normalize_brand(self._normalize(brands_row[col]) if col < len(brands_row) else None)
            if not brand or brand == "CATENA":
                skipped += 1
                continue

            total_columns += 1
            po_prefix = self._normalize(po_prefix_row[col] if col < len(po_prefix_row) else None)
            dc_code = self._normalize(dc_code_row[col] if col < len(dc_code_row) else None)
            address = self._normalize(address_row[col] if col < len(address_row) else None)
            city_state = self._normalize(city_state_row[col] if col < len(city_state_row) else None)
            zip_code = self._normalize(zip_row[col] if col < len(zip_row) else None)

            # Handle known irregular layouts (Canada/Sierra) where dc_code row contains textual entity.
            if (not dc_code) or ("WINNERS MERCHANTS" in dc_code.upper()) or ("TRADING POST" in dc_code.upper()):
                fallback_dc_code = self._normalize(alt_dc_code_row[col] if col < len(alt_dc_code_row) else None)
                if not fallback_dc_code:
                    fallback_dc_code = brand
                warnings.append(
                    f"Colonna {col+1}: dc_code anomalo/mancante ('{dc_code}'), usato fallback '{fallback_dc_code}'."
                )
                dc_code = fallback_dc_code

            dc_name = None
            if city_state and "," in city_state:
                dc_name = city_state.split(",")[0].strip().title()

            city, state, country = self._extract_city_state_country(city_state, zip_code)
            data = {
                "brand": brand,
                "dc_code": dc_code,
                "po_prefix": po_prefix,
                "name": dc_name or dc_code,
                "dc_name": dc_name,
                "address": address,
                "city": city,
                "state": state,
                "zip_code": zip_code,
                "country": country,
                "general_division_name": self._normalize(
                    general_division_row[col] if col < len(general_division_row) else None
                ),
                "general_address_line_1": self._normalize(
                    general_addr1_row[col] if col < len(general_addr1_row) else None
                ),
                "general_address_line_2": self._normalize(
                    general_addr2_row[col] if col < len(general_addr2_row) else None
                ),
                "general_address_line_3": self._normalize(
                    general_addr3_row[col] if col < len(general_addr3_row) else None
                ),
                "destination_merce_name": self._normalize(
                    dest_name_row[col] if col < len(dest_name_row) else None
                ),
                "destination_merce_dc_number": self._normalize(
                    dest_dc_row[col] if col < len(dest_dc_row) else None
                ),
                "destination_merce_address_line_1": self._normalize(
                    dest_addr1_row[col] if col < len(dest_addr1_row) else None
                ),
                "destination_merce_address_line_2": self._normalize(
                    dest_addr2_row[col] if col < len(dest_addr2_row) else None
                ),
                "destination_merce_address_line_3": self._normalize(
                    dest_addr3_row[col] if col < len(dest_addr3_row) else None
                ),
                "is_active": True,
            }

            existing = self.repo.get_by_brand_code(brand=brand, dc_code=dc_code)
            if existing is None:
                self.repo.create(data)
                inserted += 1
            else:
                for k, v in data.items():
                    setattr(existing, k, v)
                updated += 1

        self.db.commit()
        return DistributionCenterImportResponse(
            file_name=file_name,
            imported_at=datetime.now(timezone.utc),
            total_columns_read=total_columns,
            inserted_count=inserted,
            updated_count=updated,
            skipped_count=skipped,
            warnings=warnings,
        )

    def list_distribution_centers(self):
        return self.repo.list_all()

    def get_distribution_center(self, dc_id: int):
        return self.repo.get_by_id(dc_id)

    def export_distribution_centers_excel(self) -> bytes:
        rows = self.repo.list_all()
        wb = Workbook()
        ws = wb.active
        ws.title = "DISTRIBUTION_CENTERS"
        ws.append(
            [
                "BRAND",
                "DC_CODE",
                "PO_PREFIX",
                "DC_NAME",
                "ADDRESS",
                "CITY",
                "STATE",
                "ZIP",
                "COUNTRY",
                "GENERAL_DIVISION_NAME",
                "GENERAL_ADDRESS_LINE_1",
                "GENERAL_ADDRESS_LINE_2",
                "GENERAL_ADDRESS_LINE_3",
                "DESTINAZIONE_MERCE_NAME",
                "DESTINAZIONE_MERCE_DC_NUMBER",
                "DESTINAZIONE_MERCE_ADDRESS_LINE_1",
                "DESTINAZIONE_MERCE_ADDRESS_LINE_2",
                "DESTINAZIONE_MERCE_ADDRESS_LINE_3",
            ]
        )
        for r in rows:
            ws.append(
                [
                    r.brand,
                    r.dc_code,
                    r.po_prefix,
                    r.dc_name,
                    r.address,
                    r.city,
                    r.state,
                    r.zip_code,
                    r.country,
                    r.general_division_name,
                    r.general_address_line_1,
                    r.general_address_line_2,
                    r.general_address_line_3,
                    r.destination_merce_name,
                    r.destination_merce_dc_number,
                    r.destination_merce_address_line_1,
                    r.destination_merce_address_line_2,
                    r.destination_merce_address_line_3,
                ]
            )
        stream = BytesIO()
        wb.save(stream)
        return stream.getvalue()
