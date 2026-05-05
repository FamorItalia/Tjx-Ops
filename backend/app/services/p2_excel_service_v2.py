from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
import re

from openpyxl import Workbook
from openpyxl.styles import Font

from app.schemas.packing_lists import PackingListRead
from app.services.p2_pdf_service_v2 import P2RenderContext


@dataclass
class P2ExcelV2BuildResult:
    file_name: str
    file_path: str
    generated_at: datetime


class P2ExcelServiceV2:
    def __init__(self, output_root: Path):
        self.output_root = output_root

    @staticmethod
    def _slug(value: str | None, fallback: str) -> str:
        if not value:
            return fallback
        return re.sub(r"[^A-Za-z0-9]+", "_", value.upper()).strip("_") or fallback

    @staticmethod
    def _fmt_num(value: float | int | None, decimals: int = 2) -> str:
        if value is None:
            return "-"
        return f"{float(value):,.{decimals}f}".replace(",", "X").replace(".", ",").replace("X", ".")

    def build_output_path(self, payload: PackingListRead) -> Path:
        out_dir = self.output_root / "p2"
        out_dir.mkdir(parents=True, exist_ok=True)
        brand = self._slug(payload.brand, "BRAND")
        po = self._slug(payload.po_number, f"PO_{payload.id}")
        dc = self._slug(payload.dc_code, "GLOBAL")
        return out_dir / f"P2_{brand}_{po}_{dc}.xlsx"

    def build_excel(self, payload: PackingListRead, ctx: P2RenderContext) -> P2ExcelV2BuildResult:
        output_path = self.build_output_path(payload)
        wb = Workbook()
        ws = wb.active
        ws.title = "P2"

        invoice_number = (payload.invoice_number or "").strip() or "-"
        invoice_date = payload.document_date.strftime("%d/%m/%Y") if payload.document_date else "-"
        po_number = (payload.po_number or "").strip() or "-"
        po_prefix = (ctx.po_prefix or "").strip()
        po_line = f"{po_prefix} {po_number}".strip() if po_prefix else po_number

        styles_seen: set[str] = set()
        styles_ordered: list[str] = []
        for line in payload.lines:
            style = (line.vendor_style or "").strip()
            if not style:
                continue
            key = style.upper()
            if key in styles_seen:
                continue
            styles_seen.add(key)
            styles_ordered.append(style)

        ws["A1"] = "Spett."
        ws["A2"] = "EXPEDITORS INTERNATIONAL ITALIA SRL"
        ws["A3"] = "VIA SANDRO PERTINI, 56"
        ws["A5"] = "Oggetto:"
        ws["A5"].font = Font(bold=True)
        ws["A6"] = "Conferimento mandato operazione di esportazione con dest. USA con richiesta cert. P2 per paste alimentari non cotte né farcite né altrimenti preparate."
        ws["A8"] = f"Numero fattura: {invoice_number} - Data fattura: {invoice_date}"
        ws["A8"].font = Font(bold=True)
        ws["A10"] = f"PREFIX/PO {po_line}"
        ws["A11"] = "STYLE"
        ws["A12"] = ", ".join(styles_ordered) if styles_ordered else "-"
        ws["A14"] = f"NC. {ctx.customs_nc or '-'}"
        ws["A15"] = f"COLLI {self._fmt_num(payload.total_cartons, 0)}"
        ws["A16"] = f"LORDO KG {self._fmt_num(payload.total_gross_weight, 2)}"
        ws["A17"] = f"NETTO KG {self._fmt_num(payload.total_net_weight, 2)}"
        ws["A18"] = f"VALORE MERCE {self._fmt_num(ctx.valore_merce, 2)}"
        ws["A20"] = "Con la presente diamo l’incarico a codesta società ad effettuare a nome e per nostro conto, o sdoganamento della spedizione citata in oggetto con richiesta P2,"
        ws["A21"] = "e allo stesso tempo si attribuisce il potere di rappresentarci dinanzi alle autorità doganali, ai fini dell’espletamento del presente incarico."
        ws["A22"] = "Con l’occasione si porgono cordiali saluti."
        ws["A24"] = f"Livorno, {datetime.now().strftime('%d/%m/%Y')}"
        ws["A25"] = "Timbro e Firma"
        ws["A25"].font = Font(bold=True)
        ws.column_dimensions["A"].width = 180

        wb.save(output_path)
        return P2ExcelV2BuildResult(
            file_name=output_path.name,
            file_path=str(output_path.resolve()),
            generated_at=datetime.fromtimestamp(output_path.stat().st_mtime),
        )
