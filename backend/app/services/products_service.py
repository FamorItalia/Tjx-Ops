from datetime import datetime, timezone
from io import BytesIO
from pathlib import Path
import re
from uuid import uuid4

from openpyxl import Workbook, load_workbook
from sqlalchemy.orm import Session

from app.core.config import settings
from app.core.supplier_normalization import normalize_supplier_name
from app.db.models.inventory import ProductInventoryManualAdjustment
from app.repositories.products_repository import ProductsRepository
from app.schemas.products import ProductDocumentRead, ProductImportResponse, ProductRead, ProductUpdate
from app.schemas.products import ProductCreate


class ProductsService:
    def __init__(self, db: Session):
        self.db = db
        self.repo = ProductsRepository(db)

    @staticmethod
    def _style_key(value: str | None) -> str | None:
        if not value:
            return None
        key = re.sub(r"[^A-Z0-9]", "", str(value).upper())
        return key or None

    @staticmethod
    def _to_int(value) -> int | None:
        if value is None or value == "":
            return None
        try:
            return int(round(float(value)))
        except (ValueError, TypeError):
            return None

    @staticmethod
    def _to_float(value) -> float | None:
        if value is None or value == "":
            return None
        try:
            return float(value)
        except (ValueError, TypeError):
            return None

    @staticmethod
    def _to_bool(value, default: bool) -> bool:
        if value is None or value == "":
            return default
        if isinstance(value, bool):
            return value
        raw = str(value).strip().upper()
        if raw in {"1", "TRUE", "T", "YES", "Y", "SI", "S", "VERO", "X"}:
            return True
        if raw in {"0", "FALSE", "F", "NO", "N", "FALSO"}:
            return False
        return default

    @staticmethod
    def _normalize_header(value: str | None) -> str:
        if value is None:
            return ""
        return re.sub(r"\s+", " ", str(value).replace("\n", " ").strip()).upper()

    @staticmethod
    def _to_read(row) -> ProductRead:
        return ProductRead.model_validate(row, from_attributes=True)

    @staticmethod
    def _to_document_read(row) -> ProductDocumentRead:
        return ProductDocumentRead(
            id=row.id,
            product_id=row.product_id,
            file_name=row.file_name,
            content_type=row.content_type,
            size_bytes=row.size_bytes,
            uploaded_at=row.uploaded_at,
        )

    def _documents_root(self) -> Path:
        return Path(settings.templates_root) / "ANAGRAFICHE" / "PRODOTTI_DOCUMENTI"

    def _product_dir(self, product_id: int) -> Path:
        return self._documents_root() / f"PRODUCT_{product_id}"

    def import_products_excel(self, file_name: str, content: bytes) -> ProductImportResponse:
        wb = load_workbook(BytesIO(content), data_only=True)
        ws = wb[wb.sheetnames[0]]

        header_values = list(next(ws.iter_rows(min_row=1, max_row=1, values_only=True)))
        headers = [self._normalize_header(h) for h in header_values]
        header_idx = {h: i for i, h in enumerate(headers) if h}

        def col(*names: str):
            for name in names:
                idx = header_idx.get(name)
                if idx is not None:
                    return idx
            return None

        required = "TJX STYLE"
        warnings: list[str] = []
        if required not in header_idx:
            raise ValueError(f"Colonna obbligatoria mancante: {required}")

        inserted = 0
        updated = 0
        skipped = 0
        total = 0

        for row in ws.iter_rows(min_row=2, values_only=True):
            if row is None:
                skipped += 1
                continue

            tjx_style_raw = row[col("TJX STYLE")] if col("TJX STYLE") is not None else None
            tjx_style = str(tjx_style_raw).strip() if tjx_style_raw is not None else None
            style_key = self._style_key(tjx_style)
            if not style_key:
                skipped += 1
                continue

            total += 1

            data = {
                "supplier_name": normalize_supplier_name(
                    str(row[col("FORNITORE")]) if col("FORNITORE") is not None and row[col("FORNITORE")] is not None else None
                ),
                "tjx_style": tjx_style,
                "tjx_style_key": style_key,
                "description": (str(row[col("DESCRIZIONE")]).strip() if col("DESCRIZIONE") is not None and row[col("DESCRIZIONE")] is not None else None),
                "pcs_per_crt": self._to_int(row[col("PCS PER CRT")]) if col("PCS PER CRT") is not None else None,
                "strat_x_pl": self._to_int(row[col("STRAT X PL.")]) if col("STRAT X PL.") is not None else None,
                "strat_x_plt": self._to_int(row[col("STRAT X PLT")]) if col("STRAT X PLT") is not None else None,
                "layers_per_pallet": self._to_int(row[col("STRAT X PL.")]) if col("STRAT X PL.") is not None else None,
                "cartons_per_layer": self._to_int(row[col("CRTS X STRAT")]) if col("CRTS X STRAT") is not None else None,
                "cartons_per_pallet": self._to_int(row[col("CRTS PER PALLET")]) if col("CRTS PER PALLET") is not None else None,
                "carton_width_cm": self._to_float(row[col("MISURA CARTONE LARGHEZZA")]) if col("MISURA CARTONE LARGHEZZA") is not None else None,
                "carton_depth_cm": self._to_float(row[col("MISURA CARTONE PROFONDITA")]) if col("MISURA CARTONE PROFONDITA") is not None else None,
                "carton_height_cm": self._to_float(row[col("MISURA CARTONE ALTEZZA")]) if col("MISURA CARTONE ALTEZZA") is not None else None,
                "vol": self._to_float(row[col("VOL")]) if col("VOL") is not None else None,
                "peso_lordo": self._to_float(row[col("PESO LORDO")]) if col("PESO LORDO") is not None else None,
                "peso_netto": self._to_float(row[col("PESO NETTO")]) if col("PESO NETTO") is not None else None,
                "pallet_width_cm": self._to_float(row[col("LARGHEZZA PALLET")]) if col("LARGHEZZA PALLET") is not None else None,
                "pallet_depth_cm": self._to_float(row[col("PROFONDITA PALLET")]) if col("PROFONDITA PALLET") is not None else None,
                "purchase_cost_eur": (
                    self._to_float(row[col("COSTO ACQUISTO", "COSTO", "PURCHASE COST", "PURCHASE_COST")])
                    if col("COSTO ACQUISTO", "COSTO", "PURCHASE COST", "PURCHASE_COST") is not None
                    else None
                ),
                "sale_price_eur": (
                    self._to_float(row[col("PREZZO VENDITA", "VENDITA", "SALE PRICE", "SALE_PRICE")])
                    if col("PREZZO VENDITA", "VENDITA", "SALE PRICE", "SALE_PRICE") is not None
                    else None
                ),
                "document_dle": self._to_bool(
                    row[col("HAS_DLE", "DLE", "DOCUMENT_DLE")]
                    if col("HAS_DLE", "DLE", "DOCUMENT_DLE") is not None
                    else None,
                    default=True,
                ),
                "document_packing_list": self._to_bool(
                    row[col("HAS_PACKING_LIST", "PACKING_LIST", "DOCUMENT_PACKING_LIST")]
                    if col("HAS_PACKING_LIST", "PACKING_LIST", "DOCUMENT_PACKING_LIST") is not None
                    else None,
                    default=True,
                ),
                "document_sfarinati": self._to_bool(
                    row[col("HAS_SFARINATI", "SFARINATI", "DOCUMENT_SFARINATI")]
                    if col("HAS_SFARINATI", "SFARINATI", "DOCUMENT_SFARINATI") is not None
                    else None,
                    default=False,
                ),
                "document_p2": self._to_bool(
                    row[col("HAS_P2", "P2", "DOCUMENT_P2")]
                    if col("HAS_P2", "P2", "DOCUMENT_P2") is not None
                    else None,
                    default=False,
                ),
            }

            existing = self.repo.get_by_style_key(style_key)
            if existing is None:
                self.repo.create(data)
                inserted += 1
            else:
                for k, v in data.items():
                    setattr(existing, k, v)
                updated += 1

        self.db.commit()

        if "STRAT X PLT" not in header_idx:
            warnings.append("Colonna 'STRAT X PLT' non trovata nel file: campo lasciato nullo.")

        return ProductImportResponse(
            file_name=file_name,
            imported_at=datetime.now(timezone.utc),
            total_rows_read=total,
            inserted_count=inserted,
            updated_count=updated,
            skipped_count=skipped,
            warnings=warnings,
        )

    def list_products(self):
        return self.repo.list_all()

    def create_product(self, payload: ProductCreate) -> ProductRead:
        supplier_name = normalize_supplier_name(str(payload.supplier_name).strip() or None)
        tjx_style = str(payload.tjx_style).strip()
        style_key = self._style_key(tjx_style)
        if not supplier_name:
            raise ValueError("Campo supplier_name obbligatorio.")
        if not tjx_style or not style_key:
            raise ValueError("Campo tjx_style obbligatorio.")

        existing = self.repo.get_by_style_key(style_key)
        if existing is not None:
            raise ValueError(f"Prodotto giÃ  presente con TJX STYLE: {tjx_style}")

        row = self.repo.create(
            {
                "supplier_name": supplier_name,
                "tjx_style": tjx_style,
                "tjx_style_key": style_key,
                "description": payload.description.strip() if payload.description else None,
                "pcs_per_crt": payload.pcs_per_crt,
                "purchase_cost_eur": payload.purchase_cost_eur,
                "sale_price_eur": payload.sale_price_eur,
                "inventory_tracking_enabled": payload.inventory_tracking_enabled,
                "stock_product_units": payload.stock_product_units,
                "stock_packaging_units": payload.stock_packaging_units,
                "product_usage_per_unit": payload.product_usage_per_unit,
                "packaging_usage_per_unit": payload.packaging_usage_per_unit,
                "product_alert_threshold": payload.product_alert_threshold,
                "packaging_alert_threshold": payload.packaging_alert_threshold,
                "document_dle": payload.document_dle,
                "document_packing_list": payload.document_packing_list,
                "document_sfarinati": payload.document_sfarinati,
                "document_p2": payload.document_p2,
            }
        )
        self.db.commit()
        self.db.refresh(row)
        return self._to_read(row)

    def get_product(self, product_id: int):
        return self.repo.get_by_id(product_id)

    def update_product(self, product_id: int, payload: ProductUpdate) -> ProductRead | None:
        row = self.repo.get_by_id(product_id)
        if row is None:
            return None

        prev_stock_product = row.stock_product_units
        prev_stock_packaging = row.stock_packaging_units
        data = payload.model_dump(exclude_unset=True)
        for field, value in data.items():
            setattr(row, field, value)

        stock_product_changed = "stock_product_units" in data and data.get("stock_product_units") != prev_stock_product
        stock_packaging_changed = "stock_packaging_units" in data and data.get("stock_packaging_units") != prev_stock_packaging
        if stock_product_changed or stock_packaging_changed:
            note_parts: list[str] = []
            if stock_product_changed:
                note_parts.append("stock prodotto aggiornato manualmente")
            if stock_packaging_changed:
                note_parts.append("stock packaging aggiornato manualmente")
            self.db.add(
                ProductInventoryManualAdjustment(
                    product_id=row.id,
                    stock_product_before=prev_stock_product,
                    stock_product_after=row.stock_product_units,
                    stock_packaging_before=prev_stock_packaging,
                    stock_packaging_after=row.stock_packaging_units,
                    note="; ".join(note_parts),
                )
            )

        self.db.commit()
        self.db.refresh(row)
        return self._to_read(row)

    def list_product_documents(self, product_id: int) -> list[ProductDocumentRead]:
        return [self._to_document_read(x) for x in self.repo.list_documents_by_product(product_id)]

    def get_product_document(self, product_id: int, document_id: int):
        return self.repo.get_document_by_id(product_id=product_id, document_id=document_id)

    def delete_product_document(self, product_id: int, document_id: int) -> bool:
        row = self.repo.get_document_by_id(product_id=product_id, document_id=document_id)
        if row is None:
            return False
        file_path = Path(row.file_path)
        self.repo.delete_document(product_id=product_id, document_id=document_id)
        self.db.commit()
        if file_path.exists():
            try:
                file_path.unlink()
            except OSError:
                pass
        return True

    def upload_product_document(
        self,
        product_id: int,
        file_name: str,
        content: bytes,
        content_type: str | None,
    ) -> ProductDocumentRead:
        row = self.repo.get_by_id(product_id)
        if row is None:
            raise ValueError(f"Prodotto non trovato: {product_id}")

        clean_name = Path(file_name).name.strip() or f"document_{uuid4().hex}.bin"
        safe_name = re.sub(r"[^A-Za-z0-9._-]+", "_", clean_name)
        stored_name = f"{datetime.now(timezone.utc).strftime('%Y%m%d_%H%M%S')}_{uuid4().hex[:8]}_{safe_name}"

        product_dir = self._product_dir(product_id)
        product_dir.mkdir(parents=True, exist_ok=True)
        target_path = product_dir / stored_name
        target_path.write_bytes(content)

        doc = self.repo.create_document(
            {
                "product_id": product_id,
                "file_name": clean_name,
                "file_path": str(target_path),
                "content_type": content_type,
                "size_bytes": len(content),
            }
        )
        self.db.commit()
        self.db.refresh(doc)
        return self._to_document_read(doc)

    def export_products_excel(self) -> bytes:
        rows = self.repo.list_all()
        wb = Workbook()
        ws = wb.active
        ws.title = "PRODOTTI"

        headers = [
            "FORNITORE",
            "TJX STYLE",
            "DESCRIZIONE",
            "PCS PER CRT",
            "STRAT X PL.",
            "CRTS X STRAT",
            "CRTS PER PALLET",
            "MISURA CARTONE LARGHEZZA",
            "MISURA CARTONE PROFONDITA",
            "MISURA CARTONE ALTEZZA",
            "VOL",
            "STRAT X PLT",
            "PESO LORDO",
            "PESO NETTO",
            "LARGHEZZA PALLET",
            "PROFONDITA PALLET",
            "COSTO ACQUISTO",
            "PREZZO VENDITA",
            "HAS_DLE",
            "HAS_PACKING_LIST",
            "HAS_SFARINATI",
            "HAS_P2",
        ]
        ws.append(headers)

        for p in rows:
            ws.append(
                [
                    p.supplier_name,
                    p.tjx_style,
                    p.description,
                    p.pcs_per_crt,
                    p.strat_x_pl,
                    p.cartons_per_layer,
                    p.cartons_per_pallet,
                    p.carton_width_cm,
                    p.carton_depth_cm,
                    p.carton_height_cm,
                    p.vol,
                    p.strat_x_plt,
                    p.peso_lordo,
                    p.peso_netto,
                    p.pallet_width_cm,
                    p.pallet_depth_cm,
                    p.purchase_cost_eur,
                    p.sale_price_eur,
                    p.document_dle,
                    p.document_packing_list,
                    p.document_sfarinati,
                    p.document_p2,
                ]
            )

        stream = BytesIO()
        wb.save(stream)
        return stream.getvalue()
