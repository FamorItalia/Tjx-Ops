from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
import re

from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Font

from app.core.config import settings
from app.schemas.purchase_orders import PurchaseOrderPreviewResponse


@dataclass
class ExcelBuildResult:
    file_name: str
    file_path: str
    generated_at: datetime
    line_count: int


class PurchaseOrderExcelService:
    def __init__(self, output_root: Path):
        self.output_root = output_root

    @staticmethod
    def _slug(value: str | None, fallback: str) -> str:
        if not value:
            return fallback
        cleaned = re.sub(r"[^A-Za-z0-9]+", "_", value.upper()).strip("_")
        return cleaned or fallback

    def build_output_path(self, payload: PurchaseOrderPreviewResponse) -> Path:
        out_dir = self.output_root / "purchase_orders"
        out_dir.mkdir(parents=True, exist_ok=True)
        brand_part = self._slug(payload.brand, "BRAND")
        po_part = self._slug(payload.po or payload.po_raw or payload.po_normalized, f"ID_{payload.id}")
        file_name = f"ORDINE_FORNITORE_{brand_part}_{po_part}.xlsx"
        return out_dir / file_name

    def build_excel(self, payload: PurchaseOrderPreviewResponse) -> ExcelBuildResult:
        file_path = self.build_output_path(payload)
        wb = Workbook()
        ws = wb.active
        ws.title = "ORDINE_FORNITORE"

        logo_path = Path(settings.document_logo_path)
        start_row = 1
        if logo_path.exists():
            try:
                img = XLImage(str(logo_path))
                img.width = 280
                img.height = 70
                ws.add_image(img, "A1")
                start_row = 6
            except Exception:
                start_row = 1

        ws.cell(start_row, 1, "ORDINE FORNITORE").font = Font(bold=True, size=14)
        start_row += 2

        header = [
            ("Brand", payload.brand or ""),
            ("PO", payload.po or ""),
            ("Fornitore", payload.supplier or ""),
            ("Start Ship Date", str(payload.start_ship_date or "")),
            ("Cancel Ship Date", str(payload.cancel_ship_date or "")),
            ("% Adeguamento", f"{payload.adjustment_percent:.2f}%"),
        ]
        for i, (k, v) in enumerate(header, start=start_row):
            ws.cell(i, 1, k).font = Font(bold=True)
            ws.cell(i, 2, v)

        row = start_row + len(header) + 2
        columns = [
            "Vendor Style",
            "Item Code",
            "Descrizione",
            "Nest Code",
            "Qty Base",
            "%",
            "Qty Finale",
            "Units per DC",
        ]
        for c, name in enumerate(columns, start=1):
            ws.cell(row, c, name).font = Font(bold=True)
        row += 1

        for line in payload.lines:
            ws.cell(row, 1, line.vendor_style or "")
            ws.cell(row, 2, line.item_code or "")
            ws.cell(row, 3, line.description or "")
            ws.cell(row, 4, line.nest_code or "")
            ws.cell(row, 5, line.quantity_base or 0)
            ws.cell(row, 6, f"{(line.applied_percent or 0):.2f}%")
            ws.cell(row, 7, line.quantity_final or 0)
            ws.cell(row, 8, ", ".join(f"{dc}:{qty}" for dc, qty in sorted(line.units_per_dc.items())))
            row += 1

        row += 1
        ws.cell(row, 1, "Totale Righe").font = Font(bold=True)
        ws.cell(row, 2, payload.total_lines)
        row += 1
        ws.cell(row, 1, "Totale Qty Base").font = Font(bold=True)
        ws.cell(row, 2, payload.total_quantity_base)
        row += 1
        ws.cell(row, 1, "Totale Qty Finale").font = Font(bold=True)
        ws.cell(row, 2, payload.total_quantity_final)

        wb.save(file_path)
        return ExcelBuildResult(
            file_name=file_path.name,
            file_path=str(file_path.resolve()),
            generated_at=datetime.utcnow(),
            line_count=len(payload.lines),
        )
