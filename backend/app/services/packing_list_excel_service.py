from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
import re

from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Font

from app.core.config import settings
from app.schemas.packing_lists import PackingListRead


@dataclass
class PackingListExcelBuildResult:
    file_name: str
    file_path: str
    generated_at: datetime


class PackingListExcelService:
    def __init__(self, output_root: Path):
        self.output_root = output_root

    @staticmethod
    def _slug(value: str | None, fallback: str) -> str:
        if not value:
            return fallback
        return re.sub(r"[^A-Za-z0-9]+", "_", value.upper()).strip("_") or fallback

    @staticmethod
    def _normalize_po_prefix(prefix: str | None) -> str | None:
        if not prefix:
            return None
        text = re.sub(r"\s+", " ", str(prefix).strip())
        match = re.search(r"(\d+)", text)
        if match:
            return f"PO # {match.group(1)}"
        return text.replace("PO#", "PO #")

    def _format_po_display(self, po_number: str | None, po_prefix: str | None) -> str:
        po = (po_number or "").strip()
        if not po:
            return ""
        normalized_prefix = self._normalize_po_prefix(po_prefix)
        if normalized_prefix:
            return f"{normalized_prefix} {po}"
        return f"PO # {po}"

    def build_output_path(self, payload: PackingListRead) -> Path:
        out_dir = self.output_root / "packing_lists"
        out_dir.mkdir(parents=True, exist_ok=True)
        brand = self._slug(payload.brand, "BRAND")
        po = self._slug(payload.po_number, f"PL_{payload.id}")
        dc = self._slug(payload.dc_code, "NA")
        return out_dir / f"PACKING_LIST_{brand}_{po}_{dc}.xlsx"

    def build_excel(self, payload: PackingListRead) -> PackingListExcelBuildResult:
        file_path = self.build_output_path(payload)
        wb = Workbook()
        ws = wb.active
        ws.title = "PACKING_LIST"

        destination = payload.destinations[0] if payload.destinations else None
        po_display = self._format_po_display(payload.po_number, destination.po_prefix if destination else None)

        row = 1
        logo_path = Path(settings.document_logo_path)
        if logo_path.exists():
            try:
                img = XLImage(str(logo_path))
                img.width = 280
                img.height = 70
                ws.add_image(img, "A1")
                row = 6
            except Exception:
                row = 1

        ws.cell(row, 1, "PACKING LIST").font = Font(bold=True, size=14)
        row += 2

        headers = ["N. FT", "DATA FT", "DC_NAME", "PO", "DEPT"]
        values = [
            payload.invoice_number or "",
            str(payload.document_date or ""),
            (destination.dc_name if destination else "") or "",
            po_display,
            payload.dept_no or "",
        ]
        for i, h in enumerate(headers, start=1):
            ws.cell(row, i, h).font = Font(bold=True)
            ws.cell(row + 1, i, values[i - 1])
        row += 3

        ws.cell(row, 1, "RAGIONE SOCIALE FABBRICANTE").font = Font(bold=True)
        ws.cell(row, 2, payload.supplier_ragione_sociale or payload.supplier_name or "")
        row += 1
        ws.cell(row, 2, payload.supplier_address_full or "")
        row += 2

        ws.cell(row, 1, "RAGIONE SOCIALE GENERALE DIVISIONE").font = Font(bold=True)
        ws.cell(row, 2, destination.general_division_name if destination else "")
        row += 1
        ws.cell(row, 2, destination.general_address_line_1 if destination else "")
        row += 1
        ws.cell(row, 2, destination.general_address_line_2 if destination else "")
        row += 1
        ws.cell(row, 2, destination.general_address_line_3 if destination else "")
        row += 2

        ws.cell(row, 1, "LUOGO DI DESTINAZIONE DELLA MERCE").font = Font(bold=True)
        ws.cell(row, 2, destination.destination_merce_name if destination else "")
        row += 1
        ws.cell(row, 2, f"DC# {(destination.destination_merce_dc_number if destination else '') or payload.dc_code or ''}")
        row += 1
        ws.cell(row, 2, destination.destination_merce_address_line_1 if destination else "")
        row += 1
        ws.cell(row, 2, destination.destination_merce_address_line_2 if destination else "")
        row += 1
        ws.cell(row, 2, destination.destination_merce_address_line_3 if destination else "")
        row += 2

        cols = [
            "Vendor Style / Rif",
            "Description",
            "PCS x CRT",
            "Pieces",
            "Tot Cartons",
            "Carton Size",
            "Gross /Crt",
            "Net /Crt",
            "Tot Gross",
            "Tot Net",
            "CBM",
        ]
        for i, c in enumerate(cols, start=1):
            ws.cell(row, i, c).font = Font(bold=True)
        row += 1

        for group in payload.groups:
            first = True
            for p in group.products:
                ws.cell(row, 1, p.vendor_style or "")
                ws.cell(row, 2, p.description or "")
                ws.cell(row, 3, p.pcs_x_crt or "")
                ws.cell(row, 4, p.pieces or "")
                ws.cell(row, 5, group.total_cartons if first else "")
                ws.cell(row, 6, group.carton_size if first else "")
                ws.cell(row, 7, group.gross_weight_per_carton if first else "")
                ws.cell(row, 8, group.net_weight_per_carton if first else "")
                ws.cell(row, 9, group.total_gross_weight if first else "")
                ws.cell(row, 10, group.total_net_weight if first else "")
                ws.cell(row, 11, group.total_cubic_meters if first else "")
                if first:
                    for col in (7, 8, 9, 10):
                        ws.cell(row, col).number_format = "0.00"
                first = False
                row += 1

        row += 1
        ws.cell(row, 1, "Total Pieces").font = Font(bold=True)
        ws.cell(row, 2, payload.total_pieces)
        row += 1
        ws.cell(row, 1, "Total Cartons").font = Font(bold=True)
        ws.cell(row, 2, payload.total_cartons)
        row += 1
        ws.cell(row, 1, "Total Gross Weight").font = Font(bold=True)
        ws.cell(row, 2, payload.total_gross_weight)
        ws.cell(row, 2).number_format = "0.00"
        row += 1
        ws.cell(row, 1, "Total Net Weight").font = Font(bold=True)
        ws.cell(row, 2, payload.total_net_weight)
        ws.cell(row, 2).number_format = "0.00"
        row += 1
        ws.cell(row, 1, "Total Cubic Meters").font = Font(bold=True)
        ws.cell(row, 2, payload.total_cubic_meters)
        row += 2

        ws.cell(
            row,
            1,
            "Wood packaging materials have been used in the shipment and have been treated and marked in compliance with the ISPM 15 standards",
        ).font = Font(italic=True, size=9)

        wb.save(file_path)
        return PackingListExcelBuildResult(
            file_name=file_path.name,
            file_path=str(file_path.resolve()),
            generated_at=datetime.utcnow(),
        )
