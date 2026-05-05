from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
import re

from openpyxl import Workbook
from openpyxl.styles import Font

from app.schemas.packing_lists import PackingListRead


@dataclass
class DleExcelV2BuildResult:
    file_name: str
    file_path: str
    generated_at: datetime


class DleExcelServiceV2:
    def __init__(self, output_root: Path):
        self.output_root = output_root

    @staticmethod
    def _slug(value: str | None, fallback: str) -> str:
        if not value:
            return fallback
        return re.sub(r"[^A-Za-z0-9]+", "_", value.upper()).strip("_") or fallback

    def build_output_path(self, payload: PackingListRead) -> Path:
        out_dir = self.output_root / "dle"
        out_dir.mkdir(parents=True, exist_ok=True)
        brand = self._slug(payload.brand, "BRAND")
        po = self._slug(payload.po_number, f"PO_{payload.id}")
        dc = self._slug(payload.dc_code, "GLOBAL")
        return out_dir / f"DLE_{brand}_{po}_{dc}.xlsx"

    def build_excel(self, payload: PackingListRead) -> DleExcelV2BuildResult:
        output_path = self.build_output_path(payload)
        wb = Workbook()
        ws = wb.active
        ws.title = "DLE"

        ws["A1"] = "OGGETTO: DICHIARAZIONE DI LIBERA ESPORTAZIONE"
        ws["A1"].font = Font(bold=True, size=13)
        ws["A3"] = "RIFERIMENTO:"
        ws["A3"].font = Font(bold=True)
        ws["A4"] = f"Numero fattura: {payload.invoice_number or '-'} - Data fattura: {payload.document_date or '-'}"
        ws["A6"] = "Consapevoli di assumere ogni conseguente responsabilità, siamo a dichiararvi che tutto il materiale esportato:"

        bullets = [
            "- non rientra nell’elenco dei beni come da regolamento (CE) n. 428/2009 ... (Y901)",
            "- non rientra nell’elenco dei beni come da Reg.to (CE) n° 338/97 ... (Y900)",
            "- non rientra nell’elenco dei beni come da Reg.to (CE) n°1523/2007 ... (Y922)",
            "- non rientra nell’elenco dei beni come da Reg.to (CE) n° 116/2009 ... (Y903-Y905)",
            "- non rientra nell’elenco dei beni come da Reg.to (CE) n° 1236/2005 ... (Y904-Y906-Y907-Y908)",
            "- non rientra nell’elenco dei beni come da Reg.to (CE) n° 267/12 ... (Y920)",
            "- non è soggetto alle disposizioni del Reg.to (CE) n° 689/2008 ... (Y916-Y917).",
            "- Prodotto non soggetto alle disposizioni del regolamento (UE) n. 649/2012 ... (Y916)",
            "- non è soggetta a licenza di esportazione ... Reg.to CE 1005/2009 ... (Y902).",
            "- non rientra nell’elenco dei prodotti ... Reg.to (CE) n. 842/2006 ... (Y926)",
            "- merce non soggetta a sorveglianza ... (Y903)",
            "- prodotti diversi da quelli derivati dalla foca ... (Y032).",
            "- merce che non rientra negli allegati del Reg.to UE n.1332/2013 ... (Y935).",
            "- Le merci dichiarate non rientrano nel campo ... (Y909).",
            "- Le merci dichiarate non sono contemplate ... (Y927).",
            "- Prodotti e miscugli non contenenti efedrina ... 3201).",
            "- Prodotto non soggetto alle disposizioni del Reg.to (UE) n. 258/2012 ... (Y934)",
            "- Prodotto non soggetto alle disposizioni del Reg.to (CE) n.1013/2006 ... (Y923)",
            "- Merci diverse dal mercurio metallico ... (Y924)",
            "- Merci diverse ... regolamento (UE) 2024/573 ... (Y160)",
        ]
        row = 8
        for b in bullets:
            ws.cell(row=row, column=1, value=b)
            row += 1

        ws.cell(row=row + 1, column=1, value=f"Livorno, {datetime.now().strftime('%d/%m/%Y')}")
        ws.cell(row=row + 3, column=1, value="Timbro e Firma").font = Font(bold=True)
        ws.column_dimensions["A"].width = 180
        wb.save(output_path)
        return DleExcelV2BuildResult(
            file_name=output_path.name,
            file_path=str(output_path.resolve()),
            generated_at=datetime.fromtimestamp(output_path.stat().st_mtime),
        )
