from datetime import datetime, timezone
from io import BytesIO
import json
from pathlib import Path
import re
from uuid import uuid4

from openpyxl import Workbook, load_workbook
from sqlalchemy.orm import Session

from app.core.config import settings
from app.core.supplier_normalization import normalize_supplier_name
from app.repositories.suppliers_repository import SuppliersRepository
from app.schemas.suppliers import (
    SupplierCreate,
    SupplierDocumentRead,
    SupplierImportResponse,
    SupplierProductRead,
    SupplierRead,
    SupplierUpdate,
)


class SuppliersService:
    def __init__(self, db: Session):
        self.db = db
        self.repo = SuppliersRepository(db)

    @staticmethod
    def _normalize_header(value: str | None) -> str:
        if value is None:
            return ""
        return re.sub(r"\s+", " ", str(value).replace("\n", " ").strip()).upper()

    @staticmethod
    def _normalize_value(value) -> str | None:
        if value is None:
            return None
        txt = str(value).strip()
        return txt or None

    @staticmethod
    def _normalize_supplier_name(value: str | None) -> str | None:
        return normalize_supplier_name(value)

    @staticmethod
    def _supplier_key(value: str | None) -> str | None:
        if not value:
            return None
        return re.sub(r"[^A-Z0-9]", "", value.upper()) or None

    @staticmethod
    def _collect_emails(row: tuple, idxs: list[int]) -> list[str]:
        result: list[str] = []
        seen: set[str] = set()
        for idx in idxs:
            if idx >= len(row):
                continue
            raw = row[idx]
            if raw is None:
                continue
            email = str(raw).strip()
            if not email:
                continue
            low = email.lower()
            if low in seen:
                continue
            seen.add(low)
            result.append(email)
        return result

    @staticmethod
    def _to_read(row) -> SupplierRead:
        return SupplierRead(
            id=row.id,
            fornitore=row.name,
            ragione_sociale=row.ragione_sociale,
            indirizzo=row.indirizzo,
            cap=row.cap,
            citta=row.citta,
            provincia=row.provincia,
            paese=row.paese,
            telefono=row.telefono,
            persona_di_contatto=row.persona_di_contatto,
            emails=row.emails,
            email_subject_template=row.email_subject_template,
            email_order_template=row.email_order_template,
            is_active=row.is_active,
        )

    @staticmethod
    def _to_document_read(row) -> SupplierDocumentRead:
        return SupplierDocumentRead(
            id=row.id,
            supplier_id=row.supplier_id,
            file_name=row.file_name,
            content_type=row.content_type,
            size_bytes=row.size_bytes,
            uploaded_at=row.uploaded_at,
        )

    @staticmethod
    def _normalize_optional_text(value: str | None) -> str | None:
        if value is None:
            return None
        text = str(value).strip()
        return text or None

    def _documents_root(self) -> Path:
        return Path(settings.templates_root) / "ANAGRAFICHE" / "FORNITORI_DOCUMENTI"

    def _supplier_dir(self, supplier_id: int) -> Path:
        return self._documents_root() / f"SUPPLIER_{supplier_id}"

    def import_suppliers_excel(self, file_name: str, content: bytes) -> SupplierImportResponse:
        wb = load_workbook(BytesIO(content), data_only=True)
        ws = wb[wb.sheetnames[0]]

        header_values = list(next(ws.iter_rows(min_row=1, max_row=1, values_only=True)))
        headers = [self._normalize_header(h) for h in header_values]
        idx = {h: i for i, h in enumerate(headers) if h}

        required = "FORNITORE"
        if required not in idx:
            raise ValueError("Colonna obbligatoria mancante: FORNITORE")

        email_cols = [f"EMAIL{i}" for i in range(1, 7)]
        email_idxs = [idx[c] for c in email_cols if c in idx]

        warnings: list[str] = []
        missing_email_cols = [c for c in email_cols if c not in idx]
        if missing_email_cols:
            warnings.append(
                "Colonne email mancanti nel file: " + ", ".join(missing_email_cols)
            )

        inserted = 0
        updated = 0
        skipped = 0
        total = 0

        for row in ws.iter_rows(min_row=2, values_only=True):
            raw_fornitore = row[idx["FORNITORE"]] if idx["FORNITORE"] < len(row) else None
            fornitore = self._normalize_supplier_name(self._normalize_value(raw_fornitore))
            key = self._supplier_key(fornitore)
            if not fornitore or not key:
                skipped += 1
                continue

            total += 1
            emails = self._collect_emails(row, email_idxs)
            invalid_emails = [e for e in emails if "@" not in e]
            if invalid_emails:
                warnings.append(
                    f"Fornitore '{fornitore}': email con formato anomalo -> {', '.join(invalid_emails)}"
                )

            data = {
                "name": fornitore,
                "name_key": key,
                "ragione_sociale": self._normalize_value(row[idx["RAGIONE_SOCIALE"]]) if "RAGIONE_SOCIALE" in idx and idx["RAGIONE_SOCIALE"] < len(row) else None,
                "indirizzo": self._normalize_value(row[idx["INDIRIZZO"]]) if "INDIRIZZO" in idx and idx["INDIRIZZO"] < len(row) else None,
                "cap": self._normalize_value(row[idx["CAP"]]) if "CAP" in idx and idx["CAP"] < len(row) else None,
                "citta": self._normalize_value(row[idx["CITTA"]]) if "CITTA" in idx and idx["CITTA"] < len(row) else None,
                "provincia": self._normalize_value(row[idx["PROVINCIA"]]) if "PROVINCIA" in idx and idx["PROVINCIA"] < len(row) else None,
                "paese": self._normalize_value(row[idx["PAESE"]]) if "PAESE" in idx and idx["PAESE"] < len(row) else None,
                "telefono": self._normalize_value(row[idx["TELEFONO"]]) if "TELEFONO" in idx and idx["TELEFONO"] < len(row) else None,
                "persona_di_contatto": self._normalize_value(row[idx["PERSONA DI CONTATTO"]]) if "PERSONA DI CONTATTO" in idx and idx["PERSONA DI CONTATTO"] < len(row) else None,
                "emails_json": json.dumps(emails, ensure_ascii=False),
                "email_to": emails[0] if emails else None,
                "email_cc": ";".join(emails[1:]) if len(emails) > 1 else None,
                "is_active": True,
            }

            existing = self.repo.get_by_name_key(key)
            if existing is None:
                self.repo.create(data)
                inserted += 1
            else:
                for k, v in data.items():
                    setattr(existing, k, v)
                updated += 1

        self.db.commit()
        return SupplierImportResponse(
            file_name=file_name,
            imported_at=datetime.now(timezone.utc),
            total_rows_read=total,
            inserted_count=inserted,
            updated_count=updated,
            skipped_count=skipped,
            warnings=warnings,
        )

    def list_suppliers(self) -> list[SupplierRead]:
        return [self._to_read(x) for x in self.repo.list_all()]

    def create_supplier(self, payload: SupplierCreate) -> SupplierRead:
        fornitore = self._normalize_supplier_name(self._normalize_value(payload.fornitore))
        key = self._supplier_key(fornitore)
        if not fornitore or not key:
            raise ValueError("Campo fornitore obbligatorio.")

        existing = self.repo.get_by_name_key(key)
        if existing is not None:
            raise ValueError(f"Fornitore già presente: {fornitore}")

        dedup: list[str] = []
        seen: set[str] = set()
        for item in payload.emails:
            email = str(item).strip()
            if not email:
                continue
            low = email.lower()
            if low in seen:
                continue
            seen.add(low)
            dedup.append(email)

        row = self.repo.create(
            {
                "name": fornitore,
                "name_key": key,
                "ragione_sociale": self._normalize_optional_text(payload.ragione_sociale),
                "indirizzo": self._normalize_optional_text(payload.indirizzo),
                "cap": self._normalize_optional_text(payload.cap),
                "citta": self._normalize_optional_text(payload.citta),
                "provincia": self._normalize_optional_text(payload.provincia),
                "paese": self._normalize_optional_text(payload.paese),
                "telefono": self._normalize_optional_text(payload.telefono),
                "persona_di_contatto": self._normalize_optional_text(payload.persona_di_contatto),
                "emails_json": json.dumps(dedup, ensure_ascii=False),
                "email_to": dedup[0] if dedup else None,
                "email_cc": ";".join(dedup[1:]) if len(dedup) > 1 else None,
                "email_subject_template": self._normalize_optional_text(payload.email_subject_template),
                "email_order_template": self._normalize_optional_text(payload.email_order_template),
                "is_active": True,
            }
        )
        self.db.commit()
        self.db.refresh(row)
        return self._to_read(row)

    def get_supplier(self, supplier_id: int) -> SupplierRead | None:
        row = self.repo.get_by_id(supplier_id)
        if row is None:
            return None
        return self._to_read(row)

    def update_supplier(self, supplier_id: int, payload: SupplierUpdate) -> SupplierRead | None:
        row = self.repo.get_by_id(supplier_id)
        if row is None:
            return None

        update_data = payload.model_dump(exclude_unset=True)
        if "ragione_sociale" in update_data:
            row.ragione_sociale = self._normalize_optional_text(update_data.get("ragione_sociale"))
        if "indirizzo" in update_data:
            row.indirizzo = self._normalize_optional_text(update_data.get("indirizzo"))
        if "cap" in update_data:
            row.cap = self._normalize_optional_text(update_data.get("cap"))
        if "citta" in update_data:
            row.citta = self._normalize_optional_text(update_data.get("citta"))
        if "provincia" in update_data:
            row.provincia = self._normalize_optional_text(update_data.get("provincia"))
        if "paese" in update_data:
            row.paese = self._normalize_optional_text(update_data.get("paese"))
        if "telefono" in update_data:
            row.telefono = self._normalize_optional_text(update_data.get("telefono"))
        if "persona_di_contatto" in update_data:
            row.persona_di_contatto = self._normalize_optional_text(update_data.get("persona_di_contatto"))
        if "email_order_template" in update_data:
            row.email_order_template = self._normalize_optional_text(update_data.get("email_order_template"))
        if "email_subject_template" in update_data:
            row.email_subject_template = self._normalize_optional_text(update_data.get("email_subject_template"))

        if "emails" in update_data:
            emails_input = update_data.get("emails") or []
            dedup: list[str] = []
            seen: set[str] = set()
            for item in emails_input:
                email = str(item).strip()
                if not email:
                    continue
                low = email.lower()
                if low in seen:
                    continue
                seen.add(low)
                dedup.append(email)
            row.emails_json = json.dumps(dedup, ensure_ascii=False)
            row.email_to = dedup[0] if dedup else None
            row.email_cc = ";".join(dedup[1:]) if len(dedup) > 1 else None

        self.db.commit()
        self.db.refresh(row)
        return self._to_read(row)

    def list_supplier_products(self, supplier_id: int) -> list[SupplierProductRead]:
        supplier = self.repo.get_by_id(supplier_id)
        if supplier is None:
            return []

        supplier_key = self._supplier_key(self._normalize_supplier_name(supplier.name))
        if not supplier_key:
            return []

        result: list[SupplierProductRead] = []
        for product in self.repo.list_products():
            product_key = self._supplier_key(self._normalize_supplier_name(product.supplier_name))
            if product_key != supplier_key:
                continue
            result.append(
                SupplierProductRead(
                    id=product.id,
                    tjx_style=product.tjx_style,
                    description=product.description,
                    pcs_per_crt=product.pcs_per_crt,
                    purchase_cost_eur=product.purchase_cost_eur,
                    sale_price_eur=product.sale_price_eur,
                )
            )
        return result

    def list_supplier_documents(self, supplier_id: int) -> list[SupplierDocumentRead]:
        return [self._to_document_read(x) for x in self.repo.list_documents_by_supplier(supplier_id)]

    def get_supplier_document(self, supplier_id: int, document_id: int):
        return self.repo.get_document_by_id(supplier_id=supplier_id, document_id=document_id)

    def rename_supplier_document(
        self,
        supplier_id: int,
        document_id: int,
        file_name: str,
    ) -> SupplierDocumentRead | None:
        clean_name = Path(file_name).name.strip()
        if not clean_name:
            raise ValueError("Nome file non valido.")
        row = self.repo.update_document_name(
            supplier_id=supplier_id,
            document_id=document_id,
            file_name=clean_name,
        )
        if row is None:
            return None
        self.db.commit()
        self.db.refresh(row)
        return self._to_document_read(row)

    def delete_supplier_document(self, supplier_id: int, document_id: int) -> bool:
        row = self.repo.get_document_by_id(supplier_id=supplier_id, document_id=document_id)
        if row is None:
            return False
        file_path = Path(row.file_path)
        self.repo.delete_document(supplier_id=supplier_id, document_id=document_id)
        self.db.commit()
        if file_path.exists():
            try:
                file_path.unlink()
            except OSError:
                pass
        return True

    def upload_supplier_document(
        self,
        supplier_id: int,
        file_name: str,
        content: bytes,
        content_type: str | None,
    ) -> SupplierDocumentRead:
        supplier = self.repo.get_by_id(supplier_id)
        if supplier is None:
            raise ValueError(f"Fornitore non trovato: {supplier_id}")

        clean_name = Path(file_name).name.strip() or f"document_{uuid4().hex}.bin"
        safe_name = re.sub(r"[^A-Za-z0-9._-]+", "_", clean_name)
        stored_name = f"{datetime.now(timezone.utc).strftime('%Y%m%d_%H%M%S')}_{uuid4().hex[:8]}_{safe_name}"

        supplier_dir = self._supplier_dir(supplier_id)
        supplier_dir.mkdir(parents=True, exist_ok=True)
        target_path = supplier_dir / stored_name
        target_path.write_bytes(content)

        row = self.repo.create_document(
            {
                "supplier_id": supplier_id,
                "file_name": clean_name,
                "file_path": str(target_path),
                "content_type": content_type,
                "size_bytes": len(content),
            }
        )
        self.db.commit()
        self.db.refresh(row)
        return self._to_document_read(row)

    def export_suppliers_excel(self) -> bytes:
        rows = self.repo.list_all()
        wb = Workbook()
        ws = wb.active
        ws.title = "FORNITORI"
        headers = [
            "FORNITORE",
            "RAGIONE_SOCIALE",
            "INDIRIZZO",
            "CAP",
            "CITTA",
            "PROVINCIA",
            "PAESE",
            "TELEFONO",
            "PERSONA DI CONTATTO",
            "EMAIL1",
            "EMAIL2",
            "EMAIL3",
            "EMAIL4",
            "EMAIL5",
            "EMAIL6",
        ]
        ws.append(headers)

        for s in rows:
            emails = s.emails[:6]
            emails += [None] * (6 - len(emails))
            ws.append(
                [
                    s.name,
                    s.ragione_sociale,
                    s.indirizzo,
                    s.cap,
                    s.citta,
                    s.provincia,
                    s.paese,
                    s.telefono,
                    s.persona_di_contatto,
                    emails[0],
                    emails[1],
                    emails[2],
                    emails[3],
                    emails[4],
                    emails[5],
                ]
            )

        stream = BytesIO()
        wb.save(stream)
        return stream.getvalue()
